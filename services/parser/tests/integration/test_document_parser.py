import pytest


pytestmark = pytest.mark.integration


def _parse_file(filepath: str, original_filename: str | None = None, parser=None):
    from services.parser.service import ParserService
    from shared.config import ParserConfig

    parser_service = ParserService(parser or ParserConfig())
    return parser_service.parse_file(filepath, original_filename=original_filename)


def test_parse_txt_file_returns_text_blocks(test_txt_path):
    from services.parser.common.schema import TextBlock

    blocks = _parse_file(test_txt_path)

    assert blocks
    assert all(isinstance(block, TextBlock) for block in blocks)
    assert "人工智能" in "\n".join(block.text for block in blocks)


def test_parse_md_file_returns_text_and_table_blocks(tmp_path):
    from services.parser.common.schema import TableBlock, TextBlock

    md_file = tmp_path / "table.md"
    md_file.write_text(
        "# 数据库对比\n\n"
        "| 向量库 | 能力 |\n"
        "| --- | --- |\n"
        "| Qdrant | 过滤 |\n"
        "| Milvus | 分布式 |\n",
        encoding="utf-8",
    )

    blocks = _parse_file(str(md_file))

    assert isinstance(blocks[0], TextBlock)
    assert blocks[0].text == "# 数据库对比"
    assert any(isinstance(block, TableBlock) and ["Qdrant", "过滤"] in block.rows for block in blocks)


def test_parse_docx_file_returns_document_text_and_table_blocks(tmp_path):
    from docx import Document
    from services.parser.common.schema import TableBlock, TextBlock

    docx_file = tmp_path / "table.docx"
    doc = Document()
    doc.add_paragraph("数据库能力对比")
    doc.add_paragraph("首行\n续行")
    doc.add_paragraph()
    doc.add_paragraph("下一段")
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "向量库"
    table.rows[0].cells[1].text = "能力"
    table.rows[1].cells[0].text = "Qdrant"
    table.rows[1].cells[1].text = "过滤\n\n下一段"
    table.rows[2].cells[0].text = "Milvus"
    table.rows[2].cells[1].text = "分布式"
    doc.save(str(docx_file))

    blocks = _parse_file(str(docx_file))

    assert isinstance(blocks[0], TextBlock)
    assert blocks[0].text == "数据库能力对比"
    assert [block.text for block in blocks if isinstance(block, TextBlock)] == ["数据库能力对比", "首行\n续行", "下一段"]
    assert any(isinstance(block, TableBlock) and ["Qdrant", "过滤\n\n下一段"] in block.rows for block in blocks)


def test_parse_xlsx_file_returns_table_blocks(tmp_path):
    from openpyxl import Workbook
    from services.parser.common.schema import TableBlock

    xlsx_file = tmp_path / "table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据库能力"
    sheet.append(["向量库", "能力"])
    sheet.append(["Qdrant", "过滤"])
    sheet.append(["Notes", "首行\n\n下一段"])
    sheet.append(["Milvus", "分布式"])
    workbook.save(str(xlsx_file))
    workbook.close()

    blocks = _parse_file(str(xlsx_file))

    assert len(blocks) == 1
    assert isinstance(blocks[0], TableBlock)
    assert blocks[0].caption == "数据库能力"
    assert ["Qdrant", "过滤"] in blocks[0].rows
    assert ["Notes", "首行\n\n下一段"] in blocks[0].rows


def test_parse_pptx_file_returns_document_text_and_table_blocks(tmp_path):
    from pptx import Presentation
    from services.parser.common.schema import TableBlock, TextBlock

    pptx_file = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "能力概览"
    text_frame = slide.shapes.title.text_frame
    text_frame.add_paragraph()
    text_frame.add_paragraph().text = "第二段"
    presentation.slides.add_slide(presentation.slide_layouts[5]).shapes.title.text = "下一页"
    table = slide.shapes.add_table(2, 2, 0, 0, 1000, 1000).table
    table.cell(0, 0).text = "类型"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "PPT"
    table.cell(1, 1).text = "原生解析"
    presentation.save(str(pptx_file))

    blocks = _parse_file(str(pptx_file))

    assert isinstance(blocks[0], TextBlock)
    assert [(block.text, block.page) for block in blocks if isinstance(block, TextBlock)] == [
        ("能力概览", 1), ("第二段", 1), ("下一页", 2),
    ]
    assert any(isinstance(block, TableBlock) and block.page == 1 and ["PPT", "原生解析"] in block.rows for block in blocks)


def test_parse_invalid_files_raise_clear_errors(tmp_path):
    invalid_files = [
        ("fake.xlsx", "Invalid xlsx file"),
        ("fake.docx", "Invalid docx file"),
        ("fake.pdf", "Invalid pdf file"),
    ]

    for filename, message in invalid_files:
        path = tmp_path / filename
        path.write_text("invalid", encoding="utf-8")
        with pytest.raises(ValueError, match=message):
            _parse_file(str(path))


def test_mineru_content_list_returns_clean_blocks(tmp_path):
    import json
    from services.parser.common.schema import FormulaBlock, TableBlock
    from services.parser.providers.mineru.normalizer import read_content_list_blocks

    output_dir = tmp_path / "mineru-output"
    output_dir.mkdir()
    (output_dir / "demo_content_list.json").write_text(
        json.dumps([
            {"type": "text", "text": "稀疏向", "page_idx": 0},
            {"type": "text", "text": "量", "page_idx": 0},
            {
                "type": "table",
                "table_body": (
                    "<table>"
                    "<tr><td>稀疏向量</td><td>支持情况</td></tr>"
                    "<tr><td>Milvus</td><td>支持</td></tr>"
                    "</table>"
                ),
                "page_idx": 0,
            },
            {"type": "image", "img_path": "images/a.jpg", "img_caption": ["图片说明"]},
            {"type": "chart", "img_path": "images/chart.jpg", "content": "图表内容"},
            {"type": "equation", "text": "E=mc^2", "text_format": "latex", "page_idx": 0},
        ]),
        encoding="utf-8",
    )

    blocks = read_content_list_blocks(output_dir)

    assert blocks == [
        TableBlock(
            rows=[["稀疏向量", "支持情况"], ["Milvus", "支持"]],
            page=1,
        ),
        FormulaBlock("E=mc^2", page=1),
    ]

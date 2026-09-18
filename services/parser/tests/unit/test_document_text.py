from services.parser.common.text import clean_cjk_spaces, split_paragraphs


def test_clean_cjk_spaces_preserves_line_and_paragraph_boundaries():
    assert clean_cjk_spaces("中 文\n换 行\n\n段 落\r\n末 行") == "中文\n换行\n\n段落\r\n末行"


def test_text_parser_preserves_unpunctuated_paragraphs(tmp_path):
    from services.parser.documents.txt import TxtBlockParser

    path = tmp_path / "paragraphs.txt"
    path.write_text("首 行\n续 行\n\n末 段", encoding="utf-8")
    assert [block.text for block in TxtBlockParser().parse(str(path))] == ["首行\n续行", "末段"]


def test_split_paragraphs_keeps_blank_line_paragraphs_separate():
    text = "第一段说明合同履行背景。\n\n第二段说明代理制度背景。"

    chunks = split_paragraphs(text)

    assert chunks == ["第一段说明合同履行背景。", "第二段说明代理制度背景。"]


def test_text_parser_keeps_long_paragraph_whole(tmp_path):
    from services.parser.documents.txt import TxtBlockParser

    long = "很长的代理制度说明。" * 80
    path = tmp_path / "long.txt"
    path.write_text(long, encoding="utf-8")
    assert [block.text for block in TxtBlockParser().parse(str(path))] == [long]


def test_markdown_parser_keeps_headings_and_paragraphs_separate(tmp_path):
    from services.parser.documents.md import MdBlockParser

    markdown_file = tmp_path / "policy.md"
    markdown_file.write_text(
        "# 酒店制度\n\n"
        "第一段说明入住规则\n续 行\n\n"
        "## 退款政策\n\n"
        "第二段说明退款规则。\n",
        encoding="utf-8",
    )

    blocks = MdBlockParser().parse(str(markdown_file))
    chunks = [block.text for block in blocks]

    assert chunks == [
        "# 酒店制度", "第一段说明入住规则\n续行",
        "## 退款政策", "第二段说明退款规则。",
    ]
    assert [block.kind for block in blocks] == ["heading", "paragraph", "heading", "paragraph"]


def test_markdown_preserves_code_lists_and_tables_as_structures(tmp_path):
    from services.parser.common.schema import TableBlock, TextBlock
    from services.parser.documents.md import MdBlockParser

    code = "```markdown\n# Not a heading\n\n| A | B |\n| --- | --- |\n| 中 文 | 1 |\n```"
    items = "- First\n  - Nested\n- Second"
    path = tmp_path / "structures.md"
    path.write_text(
        "Heading\n=======\n\nParagraph\n" + items + "\n\n" + code
        + "\n\nName | Value\n--- | ---\nalpha | 1\n\nAfter", encoding="utf-8",
    )

    assert MdBlockParser().parse(str(path)) == [
        TextBlock("Heading\n=======", kind="heading"),
        TextBlock("Paragraph", kind="paragraph"),
        TextBlock("- First", kind="list_item"),
        TextBlock("  - Nested", kind="list_item"),
        TextBlock("- Second", kind="list_item"),
        TextBlock(code, kind="code"),
        TableBlock(rows=[["Name", "Value"], ["alpha", "1"]]),
        TextBlock("After", kind="paragraph"),
    ]


def test_markdown_preserves_installation_element_order(tmp_path):
    from services.parser.documents.md import MdBlockParser

    introduction = "# Deployment\n\nMaterials: /data/deploy\n\nScripts: /data/deploy/hacks"
    installation = "# 1 Preparation\n\n## 1.1 Install driver\n\n```sh\nchmod +x driver.run\n\n./driver.run\n```\n\nDisable Nouveau if necessary\n\n- Create config\n- Restart"
    next_section = "## 1.2 Install runtime\n\nInstall the runtime"
    path = tmp_path / "deployment.md"
    path.write_text(introduction + "\n\n" + installation + "\n\n" + next_section, encoding="utf-8")

    blocks = MdBlockParser().parse(str(path))
    assert [block.kind for block in blocks] == [
        "heading", "paragraph", "paragraph", "heading", "heading", "code",
        "paragraph", "list_item", "list_item", "heading", "paragraph",
    ]
    assert blocks[5].text == "```sh\nchmod +x driver.run\n\n./driver.run\n```"


def test_markdown_nested_code_remains_a_complete_code_element(tmp_path):
    from services.parser.documents.md import MdBlockParser

    code = "  ```python\n" + "  if ready:\n      print(value)\n" * 12 + "  ```"
    quote = "> ```python\n>     print(value)\n> ```"
    path = tmp_path / "nested.md"
    path.write_text("- Run this\n\n" + code + "\n\n" + quote + "\n", encoding="utf-8")
    blocks = MdBlockParser().parse(str(path))
    assert [(block.kind, block.text) for block in blocks] == [
        ("list_item", "- Run this"), ("code", code), ("code", quote),
    ]


def test_word_uses_native_heading_and_numbering_without_guessing(tmp_path):
    from docx import Document
    from services.parser.documents.docx import DocxBlockParser

    document = Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("Introduction")
    document.add_paragraph("First", style="List Bullet")
    document.add_paragraph("Second", style="List Number")
    document.add_paragraph("Looks like a heading").runs[0].bold = True
    path = tmp_path / "structure.docx"
    document.save(path)
    blocks = DocxBlockParser().parse(str(path))
    assert [block.kind for block in blocks] == ["heading", "paragraph", "list_item", "list_item", "paragraph"]


def test_ppt_uses_title_and_explicit_bullets(tmp_path):
    from pptx import Presentation
    from pptx.oxml.xmlchemy import OxmlElement
    from services.parser.documents.pptx import PptxBlockParser

    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[1])
    slide.shapes.title.text = "Title"
    body = slide.placeholders[1].text_frame
    body.paragraphs[0].text = "Item"
    bullet = OxmlElement("a:buChar")
    bullet.set("char", "-")
    body.paragraphs[0]._p.get_or_add_pPr().append(bullet)
    body.add_paragraph().text = "Inherited bullet"
    slide.shapes.add_textbox(0, 0, 100, 100).text_frame.text = "Other text"
    path = tmp_path / "structure.pptx"
    document.save(path)
    assert [block.kind for block in PptxBlockParser().parse(str(path))] == ["heading", "list_item", "list_item", "text"]

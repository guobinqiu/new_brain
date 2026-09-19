import threading
import time
import pytest

from services.parser.service import ParserService
from shared.config import ParserConfig


class SlowDocumentParser:
    ready = True

    def __init__(self):
        self.active = 0
        self.max_active = 0

    def parse_file(self, filepath: str, *, original_filename: str | None = None) -> list[dict]:
        self.active += 1
        self.max_active = max(self.max_active, self.active)
        time.sleep(0.05)
        self.active -= 1
        return [{"id": filepath, "content": original_filename or filepath, "metadata": {}}]


def test_parser_service_serializes_parse_file_calls(monkeypatch):
    import services.parser.service as service_mod

    monkeypatch.setattr(service_mod, "validate_pdf_file", lambda filepath: None)
    service = ParserService(ParserConfig())
    parser = SlowDocumentParser()
    service.pdf_parser = parser

    threads = [
        threading.Thread(target=service.parse_file, args=(f"file-{index}.pdf",), kwargs={"original_filename": f"file-{index}.pdf"})
        for index in range(2)
    ]

    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert parser.max_active == 1


def test_parser_service_uses_mineru_cloud_backend(monkeypatch):
    import services.parser.service as service_mod
    from shared.config import MineruCloudParserConfig

    class FakeMineruCloudDocumentParser:
        def __init__(self, config):
            self.config = config

    monkeypatch.setattr(service_mod, "MineruCloudDocumentParser", FakeMineruCloudDocumentParser)

    service = ParserService(ParserConfig(active="mineru_cloud", mineru_cloud=MineruCloudParserConfig()))

    assert isinstance(service.pdf_parser, FakeMineruCloudDocumentParser)


def test_parser_service_rejects_unknown_backend():
    with pytest.raises(ValueError, match="Unsupported parser backend"):
        ParserService(ParserConfig(active="unknown"))


def test_parser_service_routes_txt_to_local_parser_when_mineru_is_active(tmp_path, monkeypatch):
    import services.parser.service as service_mod

    class FailingMineruParser:
        ready = True

        def __init__(self, config):
            pass

        def start(self):
            pass

        def stop(self):
            pass

        def parse_file(self, filepath, *, original_filename=None):
            raise AssertionError("txt should not be parsed by mineru")

    monkeypatch.setattr(service_mod, "MineruDocumentParser", FailingMineruParser)
    text_file = tmp_path / "a.txt"
    text_file.write_text("本地文本解析", encoding="utf-8")

    blocks = ParserService(ParserConfig(active="mineru")).parse_file(str(text_file))

    assert blocks[0].text == "本地文本解析"


def test_parser_service_routes_pptx_to_local_parser_when_mineru_is_active(tmp_path, monkeypatch):
    import services.parser.service as service_mod
    from services.parser.common.schema import TableBlock
    from pptx import Presentation

    class FailingMineruParser:
        ready = True

        def __init__(self, config):
            pass

        def start(self):
            pass

        def stop(self):
            pass

        def parse_file(self, filepath, *, original_filename=None):
            raise AssertionError("pptx should not be parsed by mineru")

    monkeypatch.setattr(service_mod, "MineruDocumentParser", FailingMineruParser)
    pptx_file = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "能力概览"
    table = slide.shapes.add_table(2, 2, 0, 0, 1000, 1000).table
    table.cell(0, 0).text = "类型"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "PPT"
    table.cell(1, 1).text = "原生解析"
    presentation.save(str(pptx_file))

    blocks = ParserService(ParserConfig(active="mineru")).parse_file(str(pptx_file))

    assert blocks[0].text == "能力概览"
    assert any(isinstance(block, TableBlock) and ["PPT", "原生解析"] in block.rows for block in blocks)


def test_parser_service_rejects_legacy_office_local_input(tmp_path):
    from services.parser.common.schema import TextBlock

    calls = []

    class FakeDocumentBackend:
        ready = True

        def parse_file(self, filepath, *, original_filename=None):
            calls.append((filepath, original_filename))
            return [TextBlock(f"remote:{original_filename or filepath}")]

    service = ParserService(ParserConfig())
    service.pdf_parser = FakeDocumentBackend()

    for suffix in (".doc", ".xls", ".ppt"):
        source = tmp_path / f"legacy{suffix}"
        source.write_bytes(b"legacy")

        with pytest.raises(ValueError, match="Unsupported file type"):
            service.parse_file(str(source), original_filename=f"upload{suffix}")
    assert calls == []

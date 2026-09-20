from services.parser.common.table_blocks import table_html_to_blocks


def test_logical_tables_remove_only_their_own_empty_columns():
    blocks = table_html_to_blocks(
        '<table><tr><td>Name</td><td>Value</td><td>Unit</td></tr>'
        '<tr><td>A</td><td>1</td><td>ms</td></tr>'
        '<tr><td>2. Next table</td><td></td><td></td></tr>'
        '<tr><td>Name</td><td>Value</td><td></td></tr>'
        '<tr><td>B</td><td>2</td><td></td></tr></table>'
    )

    assert blocks[0].rows == [["Name", "Value", "Unit"], ["A", "1", "ms"]]
    assert blocks[2].rows == [["Name", "Value"], ["B", "2"]]


def test_table_uses_first_single_cell_row_as_title():
    blocks = table_html_to_blocks(
        "<table>"
        "<tr><td>5. 运维复杂度对比</td><td></td><td></td></tr>"
        "<tr><td>向量库</td><td>安装难度</td><td>集群管理</td></tr>"
        "<tr><td>Qdrant</td><td>Docker</td><td>K8s</td></tr>"
        "</table>",
    )

    assert len(blocks) == 1
    assert blocks[0].caption == "5. 运维复杂度对比"
    assert blocks[0].rows == [["向量库", "安装难度", "集群管理"], ["Qdrant", "Docker", "K8s"]]


def test_xlsx_sheets_are_read_as_table_blocks(tmp_path):
    from openpyxl import Workbook
    from services.parser.documents.xlsx import xlsx_to_table_blocks

    xlsx_file = tmp_path / "table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据库能力"
    sheet.append(["向量库", "能力"])
    sheet.append(["Qdrant", "过滤"])
    sheet.append([])
    sheet.append(["组件", "状态"])
    sheet.append(["Milvus", "可选"])
    workbook.save(str(xlsx_file))

    blocks = xlsx_to_table_blocks(str(xlsx_file))

    assert len(blocks) == 2
    assert blocks[0].caption == "数据库能力"
    assert blocks[0].rows == [["向量库", "能力"], ["Qdrant", "过滤"]]
    assert blocks[1].caption == "数据库能力"
    assert blocks[1].rows == [["组件", "状态"], ["Milvus", "可选"]]


def test_xlsx_image_formula_text_leaves_empty_cell(tmp_path):
    from openpyxl import Workbook
    from services.parser.documents.xlsx import xlsx_to_table_blocks

    xlsx_file = tmp_path / "images.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Image", "Description"])
    for formula in [
        '=DISPIMG("ID_DA47872AC74A4225B9165906EC90C954",1)',
        '= _xlfn.DISPIMG("ID_2", 1)',
        '=IMAGE("https://example.com/image.png")',
        '=_xlfn.image("https://example.com/image.png")',
    ]:
        sheet.append(["Item", formula, "Keep this text"])
        sheet.cell(sheet.max_row, 2).data_type = "s"
    sheet.append(["Item", "image description", "Keep this text"])
    workbook.save(xlsx_file)
    workbook.close()

    blocks = xlsx_to_table_blocks(str(xlsx_file))

    assert blocks[0].rows == [
        ["Name", "Image", "Description"],
        ["Item", "", "Keep this text"],
        ["Item", "", "Keep this text"],
        ["Item", "", "Keep this text"],
        ["Item", "", "Keep this text"],
        ["Item", "image description", "Keep this text"],
    ]

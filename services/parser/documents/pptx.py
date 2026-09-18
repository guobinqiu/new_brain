from services.parser.common.base import BlockParser
from services.parser.common.schema import Block, TextBlock
from services.parser.common.table_blocks import table_rows_to_blocks
from services.parser.common.validation import validate_pptx_file


class PptxBlockParser(BlockParser):
    def parse(self, filepath: str) -> list[Block]:
        validate_pptx_file(filepath)
        return self._read_blocks(filepath)

    def _read_blocks(self, filepath: str) -> list[Block]:
        from pptx import Presentation

        blocks: list[Block] = []
        presentation = Presentation(filepath)
        for page, slide in enumerate(presentation.slides, start=1):
            for shape in slide.shapes:
                if shape.has_table:
                    rows = [
                        [cell.text.strip() for cell in row.cells]
                        for row in shape.table.rows
                    ]
                    blocks.extend(table_rows_to_blocks("", rows, page=page))
                    continue
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            blocks.append(TextBlock(text, page=page, kind=_paragraph_kind(slide, shape, paragraph)))
        return blocks


def _paragraph_kind(slide, shape, paragraph) -> str:
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.oxml.ns import qn

    if slide.shapes.title is not None and slide.shapes.title.shape_id == shape.shape_id:
        return "heading"
    level = paragraph.level + 1
    properties = list(paragraph._p.xpath("./a:pPr"))
    properties.extend(shape._element.xpath(f"./p:txBody/a:lstStyle/a:lvl{level}pPr"))
    body_placeholder = shape.is_placeholder and shape.placeholder_format.type in {
        PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE,
    }
    if shape.is_placeholder:
        for placeholder in slide.slide_layout.placeholders:
            if placeholder.placeholder_format.idx == shape.placeholder_format.idx:
                properties.extend(placeholder._element.xpath(f"./p:txBody/a:lstStyle/a:lvl{level}pPr"))
                break
        style = "bodyStyle" if body_placeholder else "otherStyle"
        properties.extend(slide.slide_layout.slide_master._element.xpath(f"./p:txStyles/p:{style}/a:lvl{level}pPr"))
    for prop in properties:
        if prop.find(qn("a:buNone")) is not None:
            return "paragraph"
        if any(prop.find(qn(tag)) is not None for tag in ("a:buChar", "a:buAutoNum", "a:buBlip")):
            return "list_item"
    if body_placeholder:
        return "paragraph"
    return "text"

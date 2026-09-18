from contextlib import closing
import re

from services.parser.common.base import BlockParser
from services.parser.common.table_blocks import table_rows_to_blocks
from services.parser.common.schema import Block
from services.parser.common.validation import validate_xlsx_file


class XlsxBlockParser(BlockParser):
    def parse(self, filepath: str) -> list[Block]:
        validate_xlsx_file(filepath)
        return xlsx_to_table_blocks(filepath)


def xlsx_to_table_blocks(filepath: str) -> list[Block]:
    from openpyxl import load_workbook

    blocks: list[Block] = []
    with closing(load_workbook(filepath, data_only=True, read_only=True)) as workbook:
        for sheet in workbook.worksheets:
            rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            for table_rows in _split_non_empty_regions(rows):
                blocks.extend(table_rows_to_blocks(sheet.title, table_rows))
    return blocks


def _split_non_empty_regions(rows: list[list[str]]) -> list[list[list[str]]]:
    regions: list[list[list[str]]] = []
    current: list[list[str]] = []
    for row in rows:
        if any(cell for cell in row):
            current.append(_trim_empty_tail(row))
            continue
        if current:
            regions.append(current)
            current = []
    if current:
        regions.append(current)
    return regions


def _cell_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if re.fullmatch(r"=\s*(?:_xlfn\.)?(?:DISPIMG|IMAGE)\s*\(.*\)", text, flags=re.IGNORECASE | re.DOTALL):
        return ""
    return text


def _trim_empty_tail(row: list[str]) -> list[str]:
    end = len(row)
    while end > 0 and not row[end - 1]:
        end -= 1
    return row[:end]

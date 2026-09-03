from __future__ import annotations

from typing import Any

from rag.parser.common.schema import Block
from rag.parser.common.table_transform import _table_rows_to_blocks


def xlsx_to_table_blocks(filepath: str, parser_config: Any) -> list[Block]:
    from openpyxl import load_workbook

    workbook = load_workbook(filepath, data_only=True, read_only=True)
    blocks: list[Block] = []
    try:
        for sheet in workbook.worksheets:
            rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            for table_rows in _split_non_empty_regions(rows):
                blocks.extend(_table_rows_to_blocks(f"工作表：{sheet.title}", table_rows, parser_config))
    finally:
        workbook.close()
    return blocks


def _split_non_empty_regions(rows: list[list[str]]) -> list[list[list[str]]]:
    regions: list[list[list[str]]] = []
    current: list[list[str]] = []
    for row in rows:
        if any(cell for cell in row):
            current.append(_trim_empty_tail(row))
            continue
        if current:
            regions.append(current)
            current = []
    if current:
        regions.append(current)
    return regions


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _trim_empty_tail(row: list[str]) -> list[str]:
    end = len(row)
    while end > 0 and not row[end - 1]:
        end -= 1
    return row[:end]
